"""
sync_workbook.py
================
Write newly scraped polls back into MidtermPolls2026.xlsx, in the workbook's own
schema, so the spreadsheet stays the single source of truth.

    python3 sync_workbook.py              # report what is new, write a paste-ready CSV
    python3 sync_workbook.py --append     # also append the rows to a COPY of the workbook
    python3 sync_workbook.py --in-place   # append to the workbook itself (makes a backup)

The problem this solves
-----------------------
The top-up scrape in script 03 adds polls to the model, but the workbook never
sees them, so the spreadsheet falls further behind the forecast every week and
the gap has to be closed by hand. This closes it the other way: whatever the
scrape found that the workbook lacks is emitted in the workbook's exact column
order, ready to paste, or appended directly.

Why the default only writes a CSV
---------------------------------
The workbook has 130 sheets, formulas in Races and Ratings, and generated view
tabs. openpyxl can read and write it, but round-tripping a file that complex can
quietly drop things it does not model: charts, some conditional formatting, data
validation, and anything added by Excel features it does not support. Formulas
survive, but "mostly survives" is not good enough for a file that is the
project's source of truth and took real work to build. So the safe path is the
default and the destructive path is opt-in, with a timestamped backup.

Every row is written with include_in_model = "Review"
----------------------------------------------------
Not "Y". These rows were parsed by a scraper, and the whole reason the workbook
beats the scrapers is that a human decided what belongs in a forecast. Marking
them "Review" puts them in front of that judgement instead of around it: they
are visible in the sheet, excluded from the model until promoted, and the qa_flag
says where each one came from. Flip to Y once you have looked.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import re
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd

try:
    PROJECT = Path(__file__).resolve().parent
except NameError:
    PROJECT = Path.cwd()
sys.path.insert(0, str(PROJECT))

import config  # noqa: E402
import workbook as wb  # noqa: E402
from utils import get_logger  # noqa: E402

log = get_logger("sync_wb")

SHEET = "Polls"
COLS = ["poll_id", "race_id", "pollster", "partisan_sponsor", "field_start", "field_end",
        "date_posted", "population", "sample_n", "moe", "mode", "version", "version_note",
        "opt_a_label", "opt_a_party", "opt_a_pct", "opt_b_label", "opt_b_party", "opt_b_pct",
        "opt_c_label", "opt_c_party", "opt_c_pct", "other_undecided_pct", "high_frequency",
        "include_in_model", "qa_flag", "source_url", "entered_on", "margin_a_minus_b"]

POP_BACK = {"lv": "LV", "rv": "RV", "a": "A", "unknown": None}


def to_workbook_id(model_race_id: str, races: pd.DataFrame) -> str | None:
    """Model race_id -> the workbook's own id, by inverting workbook.race_id_map."""
    inv = {v: k for k, v in wb.race_id_map(races).items()}
    return inv.get(model_race_id)


def make_poll_id(race_wb_id: str, pollster: str, end: pd.Timestamp, d: float, r: float) -> str:
    """Mirror the workbook's own id shape: PREFIX-YYYYMMDD-POLLSTER-hash.

    The hash is over the fields that identify the survey, so re-running this
    script on the same poll produces the same id and the row cannot be added
    twice.
    """
    prefix = "GB" if race_wb_id == "US-GB" else race_wb_id
    slug = re.sub(r"[^A-Z0-9]", "", str(pollster).upper())[:14] or "UNKNOWN"
    key = f"{race_wb_id}|{pollster}|{end:%Y-%m-%d}|{d}|{r}"
    h = hashlib.sha1(key.encode()).hexdigest()[:4]
    return f"{prefix}-{end:%Y%m%d}-{slug}-{h}"


def party_of(name: str, fallback: str) -> str:
    s = str(name or "")
    if re.search(r"\(\s*I\s*\)", s):
        return "I"
    return fallback


def build_new_rows(path: Path) -> pd.DataFrame:
    """Polls present in the model but absent from the workbook."""
    polls = pd.read_parquet(config.DATA_PROCESSED / "polls_2026.parquet")
    polls["end_date"] = pd.to_datetime(polls["end_date"])
    scraped = polls[~polls["source"].astype(str).str.startswith("workbook")].copy()
    if not len(scraped):
        return pd.DataFrame(columns=COLS)

    races = wb._read(path, wb.SHEET_RACES, wb.REQUIRED_RACE_COLS)
    existing = wb._read(path, SHEET, wb.REQUIRED_POLL_COLS)
    existing["field_end"] = pd.to_datetime(existing["field_end"], errors="coerce")
    inv = {v: k for k, v in wb.race_id_map(races).items()}

    # Identity for de-duplication: race, normalised pollster, field end date.
    # Deliberately not the percentages, so a row already in the workbook with a
    # one-point difference (a different published version of the same survey) is
    # still recognised as the same poll and not added again.
    def key(rid, pollster, end):
        return (str(rid), wb.normalise_pollster(pollster).lower(), pd.Timestamp(end).date())

    have = {key(r.race_id, r.pollster, r.field_end)
            for r in existing.itertuples() if pd.notna(r.field_end)}

    rows, skipped = [], 0
    today = dt.date.today()
    for p in scraped.itertuples():
        wb_id = inv.get(p.race_id)
        if wb_id is None:
            skipped += 1
            continue
        if key(wb_id, p.pollster, p.end_date) in have:
            skipped += 1
            continue
        d, r = float(p.dem_pct), float(p.rep_pct)
        a_party = "I" if bool(getattr(p, "main_is_independent", False)) else "D"
        rows.append({
            "poll_id": make_poll_id(wb_id, p.pollster, p.end_date, d, r),
            "race_id": wb_id,
            "pollster": p.pollster,
            "partisan_sponsor": (str(p.partisan) if str(p.partisan) not in ("", "nan") else None),
            "field_start": (pd.to_datetime(p.start_date).date()
                            if pd.notna(p.start_date) else p.end_date.date()),
            "field_end": p.end_date.date(),
            "date_posted": None,
            "population": POP_BACK.get(str(p.population), None),
            "sample_n": (float(p.sample_size) if pd.notna(p.sample_size) else None),
            "moe": None,
            "mode": None,
            "version": 1,
            "version_note": None,
            "opt_a_label": (str(p.dem_candidate) or None),
            "opt_a_party": party_of(p.dem_candidate, a_party),
            "opt_a_pct": round(d / 100.0, 4),
            "opt_b_label": (str(p.rep_candidate) or None),
            "opt_b_party": "R",
            "opt_b_pct": round(r / 100.0, 4),
            "opt_c_label": None, "opt_c_party": None, "opt_c_pct": None,
            "other_undecided_pct": None,
            "high_frequency": None,
            "include_in_model": "Review",
            "qa_flag": ("Added by sync_workbook.py from the Wikipedia top-up scrape; base "
                        "question (not the leaners-pushed version); verify against the release "
                        "and set include_in_model to Y to use it"),
            "source_url": (str(p.source) if str(p.source) not in ("", "nan") else None),
            "entered_on": today,
            "margin_a_minus_b": round((d - r) / 100.0, 4),
        })

    out = pd.DataFrame(rows, columns=COLS)
    out = out.drop_duplicates(subset=["poll_id"])
    log.info("%d scraped poll(s) in the model, %d already in the workbook or unmappable, "
             "%d new", len(scraped), skipped, len(out))
    return out


def append_to_workbook(path: Path, rows: pd.DataFrame, dest: Path) -> None:
    """Append rows to the Polls sheet, preserving everything openpyxl can."""
    from openpyxl import load_workbook
    if dest != path:
        shutil.copy2(path, dest)
    book = load_workbook(dest)          # data_only=False keeps formulas as formulas
    ws = book[SHEET]
    header = [c.value for c in ws[1]]
    if header[:len(COLS)] != COLS:
        raise SystemExit(f"the Polls sheet's header does not match the expected column order.\n"
                         f"  expected: {COLS}\n  found:    {header}\n"
                         f"Not writing; the schema has changed.")
    # NOT ws.max_row. The Polls sheet carries about a thousand pre-formatted
    # empty rows below the data, ready for hand entry, and openpyxl counts them.
    # Appending after those would leave a gap of blank rows in the middle of the
    # table, which pandas then reads back as a thousand all-NaN polls. Scan up
    # for the last row that actually holds a poll_id.
    start = 2
    for r in range(ws.max_row, 1, -1):
        v = ws.cell(row=r, column=1).value
        if v not in (None, ""):
            start = r + 1
            break
    for i, rec in enumerate(rows.to_dict("records")):
        for j, col in enumerate(COLS, start=1):
            v = rec[col]
            if isinstance(v, float) and np.isnan(v):
                v = None
            ws.cell(row=start + i, column=j, value=v)
    book.save(dest)
    log.info("appended %d row(s) to %s, starting at row %d", len(rows), dest.name, start)


def main() -> None:
    argv = sys.argv[1:]
    path = wb.find_workbook()
    if path is None:
        raise SystemExit("No workbook found. See config.POLL_WORKBOOK.")
    log.info("workbook: %s", path)

    rows = build_new_rows(path)
    if not len(rows):
        print("\nNothing new: every scraped poll is already in the workbook.")
        return

    csv = config.DATA_MANUAL / f"new_polls_{dt.date.today():%Y%m%d}.csv"
    rows.to_csv(csv, index=False)

    print(f"\n{len(rows)} poll(s) to add:\n")
    print(f"{'race':14s} {'field end':11s} {'pollster':32s} {'A':>6s} {'B':>6s} {'margin':>7s}")
    for r in rows.sort_values(["race_id", "field_end"]).itertuples():
        print(f"{r.race_id:14s} {str(r.field_end):11s} {str(r.pollster)[:31]:32s} "
              f"{r.opt_a_pct*100:5.1f}% {r.opt_b_pct*100:5.1f}% {r.margin_a_minus_b*100:+6.1f}")
    print(f"\nPaste-ready, in the workbook's column order:\n    {csv}")
    print("Every row is marked include_in_model = Review, so none of them reach the "
          "forecast until you set it to Y.")

    if "--in-place" in argv:
        bak = path.with_name(f"{path.stem}.backup-{dt.datetime.now():%Y%m%d-%H%M%S}{path.suffix}")
        shutil.copy2(path, bak)
        log.info("backup written: %s", bak)
        append_to_workbook(path, rows, path)
        print(f"\nAppended in place. Backup: {bak}")
        print("Check the file opens cleanly before deleting the backup: openpyxl can drop "
              "charts and some formatting it does not model.")
    elif "--append" in argv:
        dest = path.with_name(f"{path.stem}.updated{path.suffix}")
        append_to_workbook(path, rows, dest)
        print(f"\nWrote a copy with the rows appended: {dest}")
        print("Review it, then replace your working file if it looks right.")


if __name__ == "__main__":
    main()
